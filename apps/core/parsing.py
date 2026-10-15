from __future__ import annotations

from pathlib import Path

from django.core.exceptions import ValidationError


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_BYTES = 5 * 1024 * 1024

PDF_MAGIC = b"%PDF"
DOCX_MAGIC = b"PK"
MAX_EXTRACT_CHARS = 80_000


def validate_resume_file(uploaded) -> str:
    name = getattr(uploaded, "name", "") or "resume"
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValidationError("Upload a PDF, DOCX, or TXT resume (max 5 MB).")

    size = getattr(uploaded, "size", 0) or 0
    if size > MAX_BYTES:
        raise ValidationError("Resume files must be 5 MB or smaller.")

    pos = uploaded.tell() if hasattr(uploaded, "tell") else 0
    header = uploaded.read(8)
    if hasattr(uploaded, "seek"):
        uploaded.seek(pos)

    if suffix == ".pdf" and not header.startswith(PDF_MAGIC):
        raise ValidationError("That PDF looks corrupted or is not a real PDF.")
    if suffix == ".docx" and not header.startswith(DOCX_MAGIC):
        raise ValidationError("That Word file does not look like a valid .docx.")
    return suffix


def extract_text_from_upload(uploaded, suffix: str) -> str:
    if hasattr(uploaded, "seek"):
        uploaded.seek(0)
    data = uploaded.read()
    if hasattr(uploaded, "seek"):
        uploaded.seek(0)
    return extract_text_from_bytes(data, suffix)


def extract_text_from_bytes(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".pdf":
        text = _from_pdf(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    cleaned = _normalize_whitespace(text)
    if len(cleaned) < 80:
        raise ValidationError(
            "Could not read enough text from that file. Try a text-based PDF or DOCX."
        )
    return cleaned[:MAX_EXTRACT_CHARS]


def _from_pdf(data: bytes) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if getattr(reader, "is_encrypted", False):
        raise ValidationError("Encrypted PDFs cannot be analyzed. Export an unlocked copy.")
    pages = []
    for page in reader.pages[:20]:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _from_docx(data: bytes) -> str:
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(data))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _normalize_whitespace(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    compact: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                compact.append("")
            blank = True
            continue
        compact.append(line)
        blank = False
    return "\n".join(compact).strip()
